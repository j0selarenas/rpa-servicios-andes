"""Lectura del Excel (bronce) y validaciones de negocio (silver).

Una fila con errores no detiene el proceso: va a cuarentena con todos sus
motivos y el resto sigue. Si el archivo no tiene las columnas esperadas, eso sí
es un error de sistema y se detiene todo.
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

COLUMNS = ("ID_Solicitud", "RUT", "Nombre", "Apellido", "Email", "Monto_CLP", "Fecha_Solicitud")
DATE_FORMATS = (
    (r"\d{4}-\d{2}-\d{2}", "%Y-%m-%d"),
    (r"\d{4}/\d{2}/\d{2}", "%Y/%m/%d"),
    (r"\d{1,2}-\d{1,2}-\d{4}", "%d-%m-%Y"),
    (r"\d{1,2}/\d{1,2}/\d{4}", "%d/%m/%Y"),
)
EMAIL = re.compile(r"[a-z0-9._%+'-]+@[a-z0-9-]+(\.[a-z0-9-]+)*\.[a-z]{2,}")


class InputFileError(ValueError):
    """El Excel no cumple el formato acordado (hoja o columnas)."""


def read_excel(path: Path, sheet: str) -> list[dict]:
    """Devuelve las filas con datos, con su número de fila real en el Excel."""
    # Se abre el archivo explícitamente para soltarlo siempre (en Windows un
    # archivo abierto queda bloqueado para otros procesos).
    with open(path, "rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            return _read_rows(workbook, sheet)
        finally:
            workbook.close()


def _read_rows(workbook, sheet: str) -> list[dict]:
    if sheet not in workbook.sheetnames:
        raise InputFileError(f"El Excel no tiene la hoja '{sheet}'")
    rows = workbook[sheet].iter_rows(values_only=True)
    header = [str(v).strip() if v is not None else "" for v in next(rows, ())]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        raise InputFileError(f"Faltan columnas en el Excel: {', '.join(missing)}")
    index = {c: header.index(c) for c in COLUMNS}
    result = []
    for number, values in enumerate(rows, start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue  # fila vacía: no es una solicitud
        row = {c: values[i] if i < len(values) else None for c, i in index.items()}
        result.append({"fila": number, **row})
    return result


def as_text(value) -> str | None:
    """Representación de bronce: el valor original como texto, sin interpretarlo."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def to_bronze(rows: list[dict]) -> list[dict]:
    """Bronce: las mismas filas en formato tabular, todo como texto y sin reglas de negocio."""
    return [{"fila": row["fila"], **{c.lower(): as_text(row[c]) for c in COLUMNS}} for row in rows]


# --- Validadores. Cada uno devuelve el valor normalizado o lanza ValueError con el motivo.

def normalize_rut(value) -> str:
    text = as_text(value).strip().upper()
    match = re.fullmatch(r"(\d{1,3}(?:\.\d{3})+|\d{1,8})-?([\dK])", text)
    if not match:
        raise ValueError(f"formato de RUT inválido ({text})")
    body, dv = match.group(1).replace(".", ""), match.group(2)
    total = sum(int(d) * (2 + i % 6) for i, d in enumerate(reversed(body)))
    expected = {10: "K", 11: "0"}.get(11 - total % 11, str(11 - total % 11))
    if dv != expected:
        raise ValueError(f"dígito verificador incorrecto ({text}, debería ser {expected})")
    return f"{int(body)}-{dv}"


def normalize_email(value) -> str:
    email = str(value).strip().lower()
    if not EMAIL.fullmatch(email) or ".." in email:
        raise ValueError(f"formato de email inválido ({email})")
    return email


def parse_amount(value) -> Decimal:
    """Acepta números de Excel y texto en formato chileno: 1.250.000,50 o $ 980.000."""
    if isinstance(value, bool):
        raise ValueError("monto no numérico")
    if isinstance(value, (int, float)):
        amount = Decimal(str(value))
    else:
        text = str(value).replace("$", "").replace(" ", "").strip()
        if not re.fullmatch(r"-?(\d{1,3}(\.\d{3})+|\d+)(,\d+)?", text):
            raise ValueError(f"monto no numérico ({as_text(value)})")
        amount = Decimal(text.replace(".", "").replace(",", "."))
    if amount <= 0:
        raise ValueError(f"el monto debe ser mayor a cero ({as_text(value)})")
    return amount.quantize(Decimal("0.01"))


def parse_date(value, cutoff: date) -> date:
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    else:
        text = str(value).strip()
        fmt = next((f for pattern, f in DATE_FORMATS if re.fullmatch(pattern, text)), None)
        if fmt is None:
            raise ValueError(f"formato de fecha no reconocido ({text})")
        try:
            result = datetime.strptime(text, fmt).date()
        except ValueError:
            raise ValueError(f"la fecha no existe ({text})") from None
    if result > cutoff:
        raise ValueError(f"fecha posterior al día de proceso ({result:%d-%m-%Y})")
    return result


def _text(value) -> str:
    return " ".join(as_text(value).split())


VALIDATIONS = (
    # columna Excel, campo silver, validador, código de rechazo
    ("ID_Solicitud", "id_solicitud", _text, "ID_INVALIDO"),
    ("RUT", "rut", normalize_rut, "RUT_INVALIDO"),
    ("Nombre", "nombre", _text, "TEXTO_INVALIDO"),
    ("Apellido", "apellido", _text, "TEXTO_INVALIDO"),
    ("Email", "email", normalize_email, "EMAIL_INVALIDO"),
    ("Monto_CLP", "monto_clp", parse_amount, "MONTO_INVALIDO"),
    ("Fecha_Solicitud", "fecha_solicitud", None, "FECHA_INVALIDA"),
)


def validate(rows: list[dict], cutoff: date) -> tuple[list[dict], list[dict]]:
    """Bronce -> silver. Devuelve (válidas, rechazadas).

    Se revisan todos los campos de cada fila para informar todos sus motivos.
    Duplicados por ID: si la fila es idéntica se conserva la primera; si el
    mismo ID trae datos distintos no hay cómo saber cuál es la correcta, así
    que todas van a cuarentena para que el área de origen lo resuelva.
    """
    valid, rejected = [], []
    for row in rows:
        record, reasons = {"fila": row["fila"]}, []
        for column, field, validator, code in VALIDATIONS:
            value = row[column]
            if value is None or str(value).strip() == "":
                reasons.append(("CAMPO_VACIO", column, "campo obligatorio vacío"))
                continue
            try:
                record[field] = parse_date(value, cutoff) if field == "fecha_solicitud" else validator(value)
            except (ValueError, InvalidOperation) as exc:
                reasons.append((code, column, str(exc)))
        if reasons:
            rejected.append(_rejection(row, reasons))
        else:
            valid.append(record)

    by_id = defaultdict(list)
    for record in valid:
        by_id[record["id_solicitud"]].append(record)
    deduplicated = []
    for records in by_id.values():
        first = records[0]
        same = all({**r, "fila": 0} == {**first, "fila": 0} for r in records)
        if same:
            deduplicated.append(first)
            extra, code, detail = records[1:], "ID_DUPLICADO", f"duplicado exacto de la fila {first['fila']}"
        else:
            filas = ", ".join(str(r["fila"]) for r in records)
            extra, code, detail = records, "ID_CONFLICTO", f"el mismo ID tiene datos distintos en las filas {filas}"
        for record in extra:
            row = next(r for r in rows if r["fila"] == record["fila"])
            rejected.append(_rejection(row, [(code, "ID_Solicitud", detail)]))

    deduplicated.sort(key=lambda r: r["fila"])
    rejected.sort(key=lambda r: r["fila"])
    return deduplicated, rejected


def _rejection(row: dict, reasons: list[tuple[str, str, str]]) -> dict:
    return {
        "fila": row["fila"],
        "id_solicitud": as_text(row["ID_Solicitud"]),
        "etapa": "validacion",
        "codigos": ";".join(dict.fromkeys(code for code, _, _ in reasons)),
        "motivos": "; ".join(f"{column}: {detail}" for _, column, detail in reasons),
    }
