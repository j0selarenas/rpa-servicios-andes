"""Carga en SQLite (bronce, silver, gold) y exportación a Excel.

Cada snapshot se carga en una sola transacción: queda completo o no queda nada.
- Bronce se inserta una sola vez por snapshot y conserva la ejecución que lo cargó.
- Silver y gold se borran y recalculan para el snapshot. Por eso repetir o
  reprocesar una fecha no duplica filas.
"""
from __future__ import annotations

import os
import sqlite3
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

SCHEMA = """
CREATE TABLE IF NOT EXISTS snapshots (
    snapshot_id      TEXT PRIMARY KEY,
    fecha_extraccion TEXT NOT NULL,
    manifiesto       TEXT NOT NULL,
    capturado_en     TEXT NOT NULL,
    ultimo_run_id    TEXT NOT NULL,
    procesado_en     TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS bronce_solicitudes (
    snapshot_id     TEXT NOT NULL REFERENCES snapshots(snapshot_id),
    fila_origen     INTEGER NOT NULL,
    id_solicitud    TEXT,
    rut             TEXT,
    nombre          TEXT,
    apellido        TEXT,
    email           TEXT,
    monto_clp       TEXT,
    fecha_solicitud TEXT,
    fuente          TEXT NOT NULL,
    archivo_origen  TEXT NOT NULL,
    cargado_en      TEXT NOT NULL,
    run_id          TEXT NOT NULL,
    PRIMARY KEY (snapshot_id, fila_origen)
);
CREATE TABLE IF NOT EXISTS bronce_tasas (
    snapshot_id    TEXT NOT NULL REFERENCES snapshots(snapshot_id),
    indicador      TEXT NOT NULL,
    fecha          TEXT NOT NULL,
    valor_original TEXT NOT NULL,
    fuente         TEXT NOT NULL,
    archivo_origen TEXT NOT NULL,
    url            TEXT NOT NULL,
    cargado_en     TEXT NOT NULL,
    run_id         TEXT NOT NULL,
    PRIMARY KEY (snapshot_id, indicador, fecha)
);
CREATE TABLE IF NOT EXISTS silver_solicitudes (
    snapshot_id     TEXT NOT NULL,
    id_solicitud    TEXT NOT NULL,
    fila_origen     INTEGER NOT NULL,
    rut             TEXT NOT NULL,
    nombre          TEXT NOT NULL,
    apellido        TEXT NOT NULL,
    email           TEXT NOT NULL,
    monto_clp       REAL NOT NULL,
    fecha_solicitud TEXT NOT NULL,
    run_id          TEXT NOT NULL,
    PRIMARY KEY (snapshot_id, id_solicitud)
);
CREATE TABLE IF NOT EXISTS silver_tasas (
    snapshot_id TEXT NOT NULL,
    indicador   TEXT NOT NULL,
    fecha       TEXT NOT NULL,
    valor       REAL NOT NULL,
    run_id      TEXT NOT NULL,
    PRIMARY KEY (snapshot_id, indicador, fecha)
);
CREATE TABLE IF NOT EXISTS silver_cuarentena (
    snapshot_id    TEXT NOT NULL,
    fila_origen    INTEGER NOT NULL,
    id_solicitud   TEXT,
    etapa          TEXT NOT NULL,
    codigos        TEXT NOT NULL,
    motivos        TEXT NOT NULL,
    archivo_origen TEXT NOT NULL,
    run_id         TEXT NOT NULL,
    PRIMARY KEY (snapshot_id, fila_origen)
);
CREATE TABLE IF NOT EXISTS gold_solicitudes (
    snapshot_id       TEXT NOT NULL,
    id_solicitud      TEXT NOT NULL,
    rut               TEXT NOT NULL,
    nombre            TEXT NOT NULL,
    apellido          TEXT NOT NULL,
    email             TEXT NOT NULL,
    fecha_solicitud   TEXT NOT NULL,
    monto_clp         REAL NOT NULL,
    monto_uf          REAL NOT NULL,
    monto_usd         REAL NOT NULL,
    valor_uf          REAL NOT NULL,
    fecha_uf          TEXT NOT NULL,
    valor_usd         REAL NOT NULL,
    fecha_usd         TEXT NOT NULL,
    dias_arrastre_usd INTEGER NOT NULL,
    fila_origen       INTEGER NOT NULL,
    archivo_origen    TEXT NOT NULL,
    archivo_uf        TEXT NOT NULL,
    archivo_usd       TEXT NOT NULL,
    run_id            TEXT NOT NULL,
    procesado_en      TEXT NOT NULL,
    PRIMARY KEY (snapshot_id, id_solicitud)
);
-- Última versión de cada solicitud, considerando todos los días procesados.
-- Es lo que debe leer Excel o Power BI.
CREATE VIEW IF NOT EXISTS gold_actual AS
SELECT id_solicitud, rut, nombre, apellido, email, fecha_solicitud,
       monto_clp, monto_uf, monto_usd, valor_uf, fecha_uf, valor_usd, fecha_usd,
       dias_arrastre_usd, fila_origen, archivo_origen, archivo_uf, archivo_usd,
       snapshot_id, run_id, procesado_en
FROM (
    SELECT g.*, ROW_NUMBER() OVER (
        PARTITION BY g.id_solicitud
        ORDER BY s.fecha_extraccion DESC, s.capturado_en DESC
    ) AS version
    FROM gold_solicitudes g JOIN snapshots s USING (snapshot_id)
)
WHERE version = 1;
"""

BRONZE_COLUMNS = ("id_solicitud", "rut", "nombre", "apellido", "email", "monto_clp", "fecha_solicitud")


def connect(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.executescript(SCHEMA)
    return conn


def _insert(conn: sqlite3.Connection, table: str, rows: list[dict], verb: str = "INSERT") -> None:
    if not rows:
        return
    columns = list(rows[0])
    sql = f"{verb} INTO {table} ({', '.join(columns)}) VALUES ({', '.join('?' * len(columns))})"
    conn.executemany(sql, [[row[c] for c in columns] for row in rows])


def load_snapshot(conn, manifest: dict, manifest_file: str, layers: dict, run_id: str, now: str) -> None:
    sid = manifest["snapshot_id"]
    excel = manifest["excel"]["ruta"]
    stamp = {"snapshot_id": sid, "run_id": run_id}

    with conn:  # commit al final o rollback completo si algo falla
        conn.execute(
            "INSERT INTO snapshots VALUES (?, ?, ?, ?, ?, ?) ON CONFLICT(snapshot_id) "
            "DO UPDATE SET ultimo_run_id = excluded.ultimo_run_id, procesado_en = excluded.procesado_en",
            (sid, manifest["fecha_extraccion"], manifest_file, manifest["creado_en"], run_id, now),
        )
        _insert(conn, "bronce_solicitudes", [
            {**stamp, "fila_origen": row["fila"], **{c: row[c] for c in BRONZE_COLUMNS},
             "fuente": "excel_solicitudes", "archivo_origen": excel, "cargado_en": now}
            for row in layers["bronze"]
        ], verb="INSERT OR IGNORE")
        _insert(conn, "bronce_tasas", [
            {**stamp, "indicador": r["indicador"], "fecha": r["fecha"].isoformat(),
             "valor_original": r["valor_original"], "fuente": "sii_web",
             "archivo_origen": r["archivo"], "url": r["url"], "cargado_en": now}
            for r in layers["rates"]
        ], verb="INSERT OR IGNORE")

        for table in ("gold_solicitudes", "silver_cuarentena", "silver_solicitudes", "silver_tasas"):
            conn.execute(f"DELETE FROM {table} WHERE snapshot_id = ?", (sid,))
        _insert(conn, "silver_tasas", [
            {**stamp, "indicador": r["indicador"], "fecha": r["fecha"].isoformat(), "valor": float(r["valor"])}
            for r in layers["rates"]
        ])
        _insert(conn, "silver_solicitudes", [
            {**stamp, "id_solicitud": r["id_solicitud"], "fila_origen": r["fila"], "rut": r["rut"],
             "nombre": r["nombre"], "apellido": r["apellido"], "email": r["email"],
             "monto_clp": float(r["monto_clp"]), "fecha_solicitud": r["fecha_solicitud"].isoformat()}
            for r in layers["silver"]
        ])
        _insert(conn, "silver_cuarentena", [
            {**stamp, "fila_origen": r["fila"], "id_solicitud": r["id_solicitud"], "etapa": r["etapa"],
             "codigos": r["codigos"], "motivos": r["motivos"], "archivo_origen": excel}
            for r in layers["quarantine"]
        ])
        _insert(conn, "gold_solicitudes", [
            {**stamp, **{k: (float(v) if k.startswith(("monto_", "valor_")) else v) for k, v in g.items()},
             "archivo_origen": excel, "procesado_en": now}
            for g in layers["gold"]
        ])


# --- Exportación ------------------------------------------------------------

EXPORT_QUERIES = {
    "bronce_solicitudes": "SELECT * FROM bronce_solicitudes WHERE snapshot_id = ? ORDER BY fila_origen",
    "bronce_tasas": "SELECT * FROM bronce_tasas WHERE snapshot_id = ? ORDER BY indicador, fecha",
    "silver_solicitudes": "SELECT * FROM silver_solicitudes WHERE snapshot_id = ? ORDER BY fila_origen",
    "silver_tasas": "SELECT * FROM silver_tasas WHERE snapshot_id = ? ORDER BY indicador, fecha",
    "cuarentena": """
        SELECT c.fila_origen, c.id_solicitud, c.etapa, c.codigos, c.motivos,
               b.rut, b.email, b.monto_clp, b.fecha_solicitud, c.archivo_origen, c.snapshot_id, c.run_id
        FROM silver_cuarentena c
        JOIN bronce_solicitudes b USING (snapshot_id, fila_origen)
        WHERE c.snapshot_id = ? ORDER BY c.fila_origen""",
    "gold": "SELECT * FROM gold_solicitudes WHERE snapshot_id = ? ORDER BY fila_origen",
}
NUMBER_FORMATS = {
    "monto_clp": "#,##0.00", "monto_uf": "#,##0.0000", "monto_usd": "#,##0.00",
    "valor_uf": "#,##0.00", "valor_usd": "#,##0.00", "valor": "#,##0.00",
}


def _sheet(workbook: Workbook, title: str, cursor: sqlite3.Cursor, typed: bool) -> None:
    sheet = workbook.create_sheet(title)
    columns = [c[0] for c in cursor.description]
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in cursor:
        values = list(row)
        if typed:  # en silver y gold las fechas van como fecha de Excel
            values = [date.fromisoformat(v) if c in ("fecha", "fecha_solicitud", "fecha_uf", "fecha_usd") else v
                      for c, v in zip(columns, values)]
        sheet.append(values)
    for index, column in enumerate(columns, start=1):
        letter = sheet.cell(1, index).column_letter
        sheet.column_dimensions[letter].width = max(12, min(45, len(column) + 4))
        for cell in sheet[letter][1:]:
            if column in NUMBER_FORMATS and typed:
                cell.number_format = NUMBER_FORMATS[column]
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"


def export_run(conn, snapshot_id: str, path: Path) -> None:
    """Contenido de cada capa para este snapshot, una hoja por tabla."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, query in EXPORT_QUERIES.items():
        typed = not title.startswith("bronce") and title != "cuarentena"
        _sheet(workbook, title, conn.execute(query, (snapshot_id,)), typed)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def export_gold_actual(conn, path: Path) -> None:
    """Archivo de ruta fija para Excel / Power BI. Se reemplaza en cada ejecución."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "gold_actual", conn.execute("SELECT * FROM gold_actual ORDER BY id_solicitud"), True)
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    workbook.save(tmp)
    os.replace(tmp, path)  # falla con PermissionError si el archivo está abierto en Excel
